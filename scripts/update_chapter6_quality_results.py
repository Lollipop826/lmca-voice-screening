from pathlib import Path
import subprocess

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path('/home/luyang/.codex/attachments/ce8e4683-77ca-4cb4-991c-553c51116b16/第6章_实验与验证_LY.docx')
OUTPUT = Path('/data/luyang/lmca-share/output/第6章_实验与验证_LY_实验结果更新版.docx')
WORKING = Path('/data/luyang/lmca-share/output/docx_ly_work_0914a/第6章_实验与验证_LY_工作稿.docx')
ACCEPT_CHANGES = Path('/home/luyang/.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents/scripts/accept_tracked_changes.py')
PYTHON_BIN = Path('/home/luyang/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python')

BLUE = '2F5B9C'
LIGHT_BLUE = 'EAF0F8'
MID_BLUE = 'D9E6F2'
BORDER = '8EA9C1'
WHITE = 'FFFFFF'
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, *, size=Pt(12), bold=False, color=BLACK):
    run.font.name = 'Times New Roman'
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.highlight_color = None
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), 'Times New Roman')
    rfonts.set(qn('w:hAnsi'), 'Times New Roman')
    rfonts.set(qn('w:eastAsia'), '宋体')


def reset_paragraph(paragraph, text, style='Body Text', *, bold=False, align=None, size=Pt(12)):
    paragraph.clear()
    paragraph.style = style
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_body(doc, text):
    p = doc.add_paragraph(style='Body Text')
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    set_run_font(r, size=Pt(12))
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=Pt(16), bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style='Body Text')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=Pt(10.5), bold=True)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for tag, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{tag}'))
        if node is None:
            node = OxmlElement(f'w:{tag}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        elem = borders.find(qn(f'w:{edge}'))
        if elem is None:
            elem = OxmlElement(f'w:{edge}')
            borders.append(elem)
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), '6')
        elem.set(qn('w:space'), '0')
        elem.set(qn('w:color'), BORDER)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn('w:tblHeader'))
    if tbl_header is None:
        tbl_header = OxmlElement('w:tblHeader')
        tr_pr.append(tbl_header)
    tbl_header.set(qn('w:val'), 'true')


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn('w:cantSplit'))
    if cant_split is None:
        cant_split = OxmlElement('w:cantSplit')
        tr_pr.append(cant_split)


def keep_table_together(table):
    """Keep compact result tables on one page when they fit."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def strip_paragraph_box(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    for tag in ('pBdr', 'shd'):
        node = ppr.find(qn(f'w:{tag}'))
        if node is not None:
            ppr.remove(node)


def style_table(table, widths=None, *, total_row=False):
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for ridx, row in enumerate(table.rows):
        prevent_row_split(row)
        for cidx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths and cidx < len(widths):
                cell.width = Inches(widths[cidx])
            fill = BLUE if ridx == 0 else (LIGHT_BLUE if ridx % 2 == 0 else WHITE)
            if total_row and ridx == len(table.rows) - 1:
                fill = MID_BLUE
            set_cell_shading(cell, fill)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(
                        run,
                        size=Pt(9.5),
                        bold=(ridx == 0),
                        color=RGBColor(255, 255, 255) if ridx == 0 else BLACK,
                    )
    if len(table.rows) <= 7:
        keep_table_together(table)


def add_table(doc, headers, rows, widths, *, total_row=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.cell(0, idx).text = value
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    style_table(table, widths, total_row=total_row)
    return table


def remove_from_heading_to_end(doc, prefix):
    body = doc.element.body
    start = None
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            text = ''.join(child.itertext())
            if text.startswith(prefix):
                start = child
                break
    if start is None:
        raise RuntimeError(f'未找到章节：{prefix}')
    deleting = False
    for child in list(body.iterchildren()):
        if child is start:
            deleting = True
        if deleting and child.tag != qn('w:sectPr'):
            body.remove(child)


doc = Document(SOURCE)

# 更新原先“待补充”的描述，并清除醒目的临时高亮/批注式排版。
reset_paragraph(
    doc.paragraphs[3],
    '本章围绕语音交互链路的实时性开展验证，重点测量ASR、Agent、TTS以及端到端总延迟，并通过长期记忆和情绪分析模块的开关对比，观察模块配置对上述延迟指标及回答质量的影响。回答质量采用112条高辨别力样本进行匿名盲评。',
)
reset_paragraph(doc.paragraphs[4], '6.1.2 实验环境与数据', style='Heading 2', bold=True, size=Pt(16))
reset_paragraph(
    doc.paragraphs[6],
    '首轮延迟实验最终纳入88条有效试验，由下列11个场景构成。每个场景测试4种模块配置，每种配置重复2次，因此每个场景产生8条有效试验，合计11×8＝88条。',
)
doc.paragraphs[6].paragraph_format.first_line_indent = Pt(24)
doc.paragraphs[6].paragraph_format.left_indent = Pt(0)
doc.paragraphs[6].paragraph_format.right_indent = Pt(0)
strip_paragraph_box(doc.paragraphs[6])
reset_paragraph(doc.paragraphs[8], '每个场景包含的8条试验结构如下：', style='Body Text')
doc.paragraphs[8].paragraph_format.keep_with_next = True
doc.paragraphs[8].paragraph_format.first_line_indent = Pt(0)
for idx in (7, 8, 9):
    strip_paragraph_box(doc.paragraphs[idx])

# 先统一原有表格，再替换回答质量预留段落。
existing_widths = {
    0: [0.55, 1.25, 2.55, 0.85, 0.8],
    1: [0.9, 2.1, 1.3, 1.7],
    2: [1.8, 1.4, 1.4, 1.4],
    3: [0.85, 2.65, 0.85, 0.85, 0.8],
    4: [1.65, 0.87, 0.87, 0.87, 0.87, 0.87],
}
for idx, table in enumerate(doc.tables[:5]):
    style_table(table, existing_widths[idx], total_row=(idx in (0, 1)))

remove_from_heading_to_end(doc, '6.2.3 回答质量评估')

add_heading(doc, '6.2.3 回答质量评估（112条高辨别力样本）')
add_body(
    doc,
    '回答质量实验包含14个高辨别力语音场景，其中包括4个记忆依赖场景、2个记忆陷阱场景、6个同文本不同语气的情绪场景和2个中性控制场景。每个场景均在M1E1、M0E1、M1E0、M0E0四种配置下重复2次，共获得112条有效系统回答。',
)
add_body(
    doc,
    '评估时隐藏模块配置，将同一场景的四条回答匿名映射为A、B、C、D。使用qwen-flash作为模型评审器，对28个配对单元各独立重复评分3次，共形成84次盲评任务；每条回答获得3次1～5分评分。评分维度包括记忆相关性、连续性、事实一致性、情绪匹配、共情、校准性、边界遵守、帮助性与安全性。表中“平均分差”均为同一配对单元内开启模块减去关闭模块，正值表示开启后得分更高。',
)

add_caption(doc, '表6.3 112条高辨别力样本与盲评完成情况')
add_table(
    doc,
    ['项目', '计划量', '完成量', '说明'],
    [
        ['系统回答', '112条', '112条', '14场景×4配置×2次'],
        ['配对单元', '28组', '28组', '同场景、同输入条件下配对'],
        ['盲评任务', '84次', '84次', '每个配对单元重复3次'],
        ['单条回答评分', '3次', '3次', '1～5分量表'],
        ['失败记录', '—', '0次', '无API或格式解析失败'],
    ],
    [1.25, 0.95, 0.95, 2.85],
)

add_body(
    doc,
    '长期记忆模块的主效应在记忆依赖场景中较为明确。开启记忆后，记忆相关性、对话连续性、共情与帮助性的95%置信区间下界均高于0；事实一致性呈正向趋势，但区间下界为0。具体结果见表6.4。',
)
add_caption(doc, '表6.4 记忆依赖场景中长期记忆模块的配对评分结果')
add_table(
    doc,
    ['评价维度', '平均分差', '95%置信区间', '胜/平/负', '结果解释'],
    [
        ['记忆相关性', '+0.7083', '[0.2083, 1.2500]', '5/2/1', '稳定正向'],
        ['对话连续性', '+0.7083', '[0.1875, 1.2708]', '5/2/1', '稳定正向'],
        ['事实一致性', '+0.2500', '[0.0000, 0.5833]', '2/6/0', '正向趋势'],
        ['共情', '+0.5625', '[0.1042, 1.1250]', '5/3/0', '稳定正向'],
        ['帮助性', '+0.7708', '[0.2500, 1.3125]', '5/3/0', '稳定正向'],
    ],
    [1.15, 1.05, 1.55, 0.95, 1.30],
)
add_body(
    doc,
    '在4个记忆陷阱配对单元中，记忆相关性、连续性、事实一致性和帮助性的平均分差均为0；在4个中性控制配对单元中，各维度差异也接近0。该结果说明本轮测试中长期记忆带来的增益主要集中在确实需要调用历史信息的场景，没有观察到在无关场景中持续放大旧信息的现象。',
)

add_body(
    doc,
    '情绪模块在本轮合成语音的同文本、不同语气场景中未表现出稳定正向增益。情绪匹配和校准性的差异接近0；共情与帮助性出现小幅负向差异；边界遵守与安全性没有变化。结果见表6.5。',
)
add_caption(doc, '表6.5 情绪场景中情绪分析模块的配对评分结果')
add_table(
    doc,
    ['评价维度', '平均分差', '95%置信区间', '胜/平/负', '结果解释'],
    [
        ['情绪匹配', '-0.0833', '[-0.1806, 0.0000]', '1/7/4', '未见正向增益'],
        ['共情', '-0.1111', '[-0.2222, -0.0139]', '1/6/5', '小幅负向'],
        ['校准性', '-0.0417', '[-0.1389, 0.0417]', '2/7/3', '无稳定差异'],
        ['边界遵守', '0.0000', '[0.0000, 0.0000]', '0/12/0', '无差异'],
        ['帮助性', '-0.1250', '[-0.2361, -0.0139]', '1/6/5', '小幅负向'],
    ],
    [1.15, 1.05, 1.55, 0.95, 1.30],
)
add_body(
    doc,
    '从四配置总体首选结果看，84次盲评任务中有70次判为并列，占83.3%；M1E0、M1E1、M0E1和M0E0分别获得7、3、3和1次单独首选。总体首选以并列为主，说明除记忆依赖场景外，多数回答之间的差异较小。',
)
add_caption(doc, '表6.6 四种配置的总体首选分布（84次盲评任务）')
add_table(
    doc,
    ['首选结果', '次数', '占比', '说明'],
    [
        ['并列', '70', '83.3%', '未形成唯一首选'],
        ['M1E0', '7', '8.3%', '记忆开、情绪关'],
        ['M1E1', '3', '3.6%', '记忆开、情绪开'],
        ['M0E1', '3', '3.6%', '记忆关、情绪开'],
        ['M0E0', '1', '1.2%', '记忆关、情绪关'],
    ],
    [1.3, 0.85, 0.85, 3.0],
    total_row=False,
)

add_heading(doc, '6.3 实验结论与限制')
add_body(
    doc,
    '本轮实验完成了对ASR、Agent、TTS和端到端总延迟的统一测量。结果显示，当前链路的总体首响中位数为3418.4 ms，其中ASR最终转写约109.6 ms，Agent首段文字生成约1664.9 ms，TTS首音频启动约295.3 ms。后续降低总延迟时，应优先优化Agent首段生成和端到端等待策略。',
)
add_body(
    doc,
    '回答质量方面，长期记忆模块在记忆依赖场景中对记忆相关性、连续性、共情和帮助性产生了稳定正向作用，而在记忆陷阱与中性控制场景中差异接近0。情绪模块在本轮合成语音样本上未形成稳定正向增益，共情和帮助性还出现小幅负向差异，因此当前结果更支持继续优化情绪信号的使用方式，而不是据此认定情绪分析对所有真实交互均无效。',
)
add_body(
    doc,
    '本章延迟结果来自合成语音和环回事件链路，不包含真实麦克风、扬声器、浏览器渲染和物理播放延迟；8条未形成完整事件链路的首轮试验未纳入延迟统计，但原始记录已保留。回答质量评分由同一模型评审器重复3次完成，不等同于3名独立人工评审；情绪语音为合成样本，且缺少独立人工情绪标签。因此，表6.4～表6.6用于反映当前测试条件下的相对趋势，后续仍需通过真实语音、人工双盲评分和更多用户样本复核。',
)

# 清理全文残留高亮，并保证普通正文文字颜色统一；表头随后维持白字。
for p in doc.paragraphs:
    for run in p.runs:
        run.font.highlight_color = None
        if p.style.name not in ('Heading 1', 'Heading 2'):
            run.font.color.rgb = BLACK

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
WORKING.parent.mkdir(parents=True, exist_ok=True)
doc.save(WORKING)
subprocess.run(
    [str(PYTHON_BIN), str(ACCEPT_CHANGES), str(WORKING), '--mode', 'accept', '--out', str(OUTPUT)],
    check=True,
)
print(OUTPUT)
